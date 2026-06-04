from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def add_title_slide(prs, title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = title_text
    subtitle.text = subtitle_text
    
    # Styling
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 0, 128) # Purplle
    title.text_frame.paragraphs[0].font.bold = True

def add_bullet_slide(prs, title_text, bullet_points):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 0, 128)
    
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = bullet_points[0]
    
    for point in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title
    add_title_slide(prs, 
                   "Purplle Store Intelligence System", 
                   "Real-time retail analytics platform for physical Purplle stores")
    
    # Slide 2: Overview
    add_bullet_slide(prs, "Overview", [
        "Ingests CCTV video streams and POS transactions in real-time.",
        "Answers key questions: Who entered? Where did they go? Did they buy?",
        "Produces actionable insights via a live web dashboard.",
        "Zero-infrastructure setup, completely offline capable."
    ])
    
    # Slide 3: Architecture
    add_bullet_slide(prs, "Architecture & Flow", [
        "1. Pipeline: CCTV videos processed locally.",
        "2. Object Detection: YOLOv8 Nano for lightweight CPU inference.",
        "3. Re-ID Tracker: Cross-camera tracking using HSV color histograms.",
        "4. API Server: FastAPI ingests events & POS data into SQLite.",
        "5. Dashboard: Real-time HTML/JS dashboard polling every 5s."
    ])
    
    # Slide 4: Key Metrics
    add_bullet_slide(prs, "Key Metrics Tracked", [
        "Conversion Rate: Matches billing zone visitors to POS transactions.",
        "Average Dwell Time: Tracks duration spent per product zone.",
        "Queue Depth: Live count of customers in the billing queue.",
        "Abandonment Rate: Identifies customers leaving the queue without buying."
    ])
    
    # Slide 5: Real-time Anomaly Detection
    add_bullet_slide(prs, "Anomaly Detection Rules", [
        "Billing Queue Spikes: Alerts when queue depth > 5 (WARN) or > 10 (CRITICAL).",
        "Conversion Drops: Alerts when today's rate falls below 70% of 7-day average.",
        "Dead Zones: Flags zones with no new visits in 30+ minutes.",
        "Proactive alerts allow store managers to act immediately."
    ])
    
    # Slide 6: Live Dashboard
    add_bullet_slide(prs, "Dashboard Features", [
        "Live KPI Cards: Visitors, conversion, queue depth.",
        "Conversion Funnel: Entry → Zone Visit → Billing Queue → Purchase.",
        "Zone Heatmap: Color-coded map of store visit frequencies.",
        "Event Stream: Scrolling live feed of CCTV-derived events.",
        "Dark Mode Premium UI: Clean, responsive design."
    ])
    
    prs.save('d:/purple_hackathon/Purplle_Store_Intelligence.pptx')
    print("Presentation saved to d:/purple_hackathon/Purplle_Store_Intelligence.pptx")

if __name__ == '__main__':
    create_presentation()
